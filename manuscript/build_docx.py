"""Build manuscript_draft.docx from manuscript_draft.md with journal-style formatting.

    python manuscript/build_docx.py

Times New Roman throughout; body 12 pt justified; title 16 pt, heading 1 14 pt, heading 2
12 pt; tables and captions 10 pt; equations are native Word (OMML) equations from LaTeX.
"""

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
FONT = "Times New Roman"


def set_font(style, size, bold=None, italic=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rfonts.attrib.pop(qn(theme_attr), None)
    style.font.color.rgb = None
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic


def reference_doc(path: Path) -> None:
    data = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                          capture_output=True, check=True).stdout
    path.write_bytes(data)
    doc = Document(path)
    styles = {s.name: s for s in doc.styles}
    spec = {
        "Normal": (12, None, None), "Body Text": (12, None, None), "First Paragraph": (12, None, None),
        "Compact": (12, None, None), "Title": (16, True, False), "Subtitle": (12, False, True),
        "Author": (12, None, None), "Date": (12, None, None), "Abstract": (12, None, None),
        "Heading 1": (14, True, False), "Heading 2": (12, True, False), "Heading 3": (12, True, True),
        "Caption": (10, None, True), "Image Caption": (10, None, True), "Table Caption": (10, None, True),
        "Bibliography": (12, None, None), "Block Text": (12, None, None),
    }
    for name, (size, bold, italic) in spec.items():
        if name in styles:
            set_font(styles[name], size, bold, italic)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Abstract", "Bibliography", "Block Text"):
        if name in styles:
            pf = styles[name].paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.line_spacing = 1.15
            pf.space_after = Pt(6)
    for name in ("Title", "Subtitle", "Author", "Date"):
        if name in styles:
            styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in ("Caption", "Image Caption", "Table Caption"):
        if name in styles:
            styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.54)
    doc.save(path)


def finish(path: Path) -> None:
    """Enforce fonts and justification on every run/paragraph (pandoc sets some directly)."""
    doc = Document(path)
    centred = {"Title", "Subtitle", "Author", "Date", "Captioned Figure"}
    def fix(par, size):
        name = par.style.name
        if par._p.findall(".//" + qn("m:oMathPara")):
            pass  # display equations keep Word's centred equation layout
        elif name in centred:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif name.lower().startswith("heading"):
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in par.runs:
            run.font.name = FONT
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
            if size:
                run.font.size = Pt(size)
    for sec in doc.sections:  # A4, 2.54 cm margins
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.54)
    for par in doc.paragraphs:
        fix(par, None)
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in table.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    fix(par, 10)
    doc.save(path)


EQ_NUM = re.compile(r"\\qquad\s*\((\d+[a-z]?)\)\s*\$\$")


def strip_equation_numbers(md: str) -> tuple[str, list[str]]:
    """Remove the trailing '\\qquad (n)' from display equations; return the numbers in order."""
    numbers = EQ_NUM.findall(md)
    return EQ_NUM.sub("$$", md), numbers


def number_equations(path: Path, numbers: list[str]) -> None:
    """Place each display equation in a borderless 2-cell row: equation centred, '(n)' flush right."""
    doc = Document(path)
    paras = [p for p in doc.paragraphs if p._p.findall(".//" + qn("m:oMathPara"))]
    if len(paras) != len(numbers):
        raise ValueError(f"{len(paras)} display equations but {len(numbers)} numbers")
    width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for par, num in zip(paras, numbers):
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        tbl_pr = table._tbl.tblPr
        borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(borders.makeelement(qn(f"w:{edge}"), {qn("w:val"): "nil"}))
        tbl_pr.append(borders)
        layout = tbl_pr.makeelement(qn("w:tblLayout"), {qn("w:type"): "fixed"})
        tbl_pr.append(layout)
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = tbl_pr.makeelement(qn("w:tblW"), {})
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:w"), str(int(width.twips if hasattr(width, "twips") else width / 635)))
        tbl_w.set(qn("w:type"), "dxa")
        eq_cell, num_cell = table.rows[0].cells
        eq_cell.width, num_cell.width = width - Cm(1.6), Cm(1.6)
        table.columns[0].width, table.columns[1].width = width - Cm(1.6), Cm(1.6)
        par._p.addprevious(table._tbl)
        eq_cell.paragraphs[0]._p.addprevious(par._p)
        eq_cell._tc.remove(eq_cell.paragraphs[-1]._p)
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        num_par = num_cell.paragraphs[0]
        num_par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = num_par.add_run(f"({num})")
        run.font.name, run.font.size = FONT, Pt(12)
        num_cell.vertical_alignment = 1  # centre
    doc.save(path)


def main() -> None:
    ref = HERE / "reference.docx"
    reference_doc(ref)
    md, numbers = strip_equation_numbers((HERE / "manuscript_draft.md").read_text(encoding="utf-8"))
    tmp = HERE / ".build_tmp.md"
    tmp.write_text(md, encoding="utf-8")
    out = HERE / "manuscript_draft.docx"
    try:
        subprocess.run(["pandoc", str(tmp), "-o", str(out), "--reference-doc", str(ref),
                        "--resource-path", f"{HERE}:{HERE.parent}"], check=True)
    finally:
        tmp.unlink()
    finish(out)
    number_equations(out, numbers)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
